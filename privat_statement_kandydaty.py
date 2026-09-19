#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""privat_statement_kandydaty.py — кандидати з ЩОДЕННОЇ виписки ПриватБанку (усі рухи по
рахунку), яких НЕ видно в книзі.

НАВІЩО (запит власника 2026-09-19, "виписки з банка приходять кожен день на пошту"): лист
ПриватБанку — це НЕ сама виписка, а нагадування з кнопкою "Отримати виписку", яка веде на
підписане пряме посилання `att.privatbank.ua/efile/<токен>` (через awstrack.me-трекер +
socauth.privatbank.ua/out_click.php-редирект) — PDF, ПІДПИСАНИЙ CAdES (PKCS7-обгортка навколо
`%PDF-…`). Перевірено живо 2026-09-19: посилання скачується БЕЗ логіну в Приват24 (не той шлях,
що заблокований Автоклієнт API — `a6fbc01`, платний тариф). `kodv_mail_archiver.py` витягує це
посилання й зберігає PDF у документи_КОДВ/YYYY-MM/ПриватБанк/ — той самий патерн, що
`_find_rozetkapay_link`/`_download` для RozetkaPay (PR #566).

ЩО: читає ВСІ PDF у документи_КОДВ/*/ПриватБанк/*.pdf ЩОРАЗУ (без власного курсора — навмисно,
див. нижче), парсить таблицю операцій (`pdfplumber.extract_tables()` — НЕ `extract_text()`: для
широкої багаторядкової таблиці текстовий потік плутає колонки, таблична екстракція дає чисті
клітинки). Кожну операцію звіряє з книгою за (сума, дата ±1 день) — той самий принцип, що
`checkbox_registry_sync._match_book`, лише ширше: шукає суму НЕ в одній конкретній графі
(виписка покриває дохід/комісії/інші списання одразу), а в БУДЬ-ЯКІЙ графі СИРИХ операцій книги
(2,3,6,7,8,9,10 — БЕЗ графи 4 й графи 11, це ОБЧИСЛЮВАНІ підсумки, не окремі операції, див.
`_BOOK_MONEY_COLS`) на відповідну дату — бо яка саме графа підходить конкретній операції,
вирішує бухгалтер, не скрипт. Пише кандидатів у документи_КОДВ, книгу НЕ чіпає.

ЧОМУ БЕЗ ВЛАСНОГО КУРСОРА (на відміну від novapay_registry_kandydaty.py): цей скрипт передає
"поточний повний список ще-не-в-книзі" у `kandydaty_registry.sync_open_candidates(resolve=True)`
(патерн `toysi_returns_kandydaty.py`/`vchasno_akty_kandydaty.py`) — реєстр сам закриває кандидата,
щойно він зникає зі списку (бо `_already_in_book` тепер бачить його в книзі). Якби скрипт мав
власний курсор "вже показував — не показуй знову", той самий кандидат випав би зі списку
`unresolved` на ДРУГОМУУ прогоні незалежно від того, чи бухгалтер справді його внесла — і
`resolve=True` хибно закрив би його як "resolved" (рецидив класу бага з КОДВ_журнал
«ДОПОВНЕННЯ 5», Д1). PDF повторно парситься щоразу — дешево (кілька файлів, десятки KB).

Кирилична деталізація операцій (призначення платежу, найменування контрагента) з PDF
екстрактиться КОРЕКТНО (перевірено живо — Windows-консоль просто не вміє її друкувати в
терміналі, це НЕ дефект PDF/шрифту) — але для звірки з книгою вона й не потрібна: сума+дата
достатньо, той самий принцип, що Checkbox.

ЗАПУСК: python privat_statement_kandydaty.py
"""
import glob
import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import kandydaty_registry
import source_freshness

BASE_DIR = Path(__file__).parent
COWORK_DIR = Path(os.environ.get("PLUTUS_COWORK_DIR",
                                 r"C:\Users\smach\Claude\Projects\PlutusToys_avtonomiya"))
DOCS_DIR = COWORK_DIR / "документи_КОДВ"
KODV_XLSX = COWORK_DIR / "KODV_PlutusToys_2026.xlsx"
# Грошові графи книги (0-indexed, ряд iter_rows(min_row=7)) — заголовки звірено ЖИВО з
# KODV_PlutusToys_2026.xlsx (рядок 6), не з пам'яті: 1=Графа2 сума доходу за день, 2=Графа3
# повернення коштів, 5=Графа6 витрати на товари, 6=Графа7 оплата праці/ЦПХ, 7=Графа8 ЄСВ/
# податки, 8=Графа9 інші витрати, 9=Графа10 амортизація — усе СИРІ значення операцій.
# СВІДОМО ВИКЛЮЧЕНО (не грошові кандидати для звірки, хоч і числові): 0=Графа1 дата (ключ
# звірки, не сума), 3=Графа4 "дохід за вирахуванням повернень" — ОБЧИСЛЮВАНА (графа2-графа3,
# не сира операція), 4=Графа5 "реквізити підтвердного документа" (текст), 10=Графа11 "чистий
# оподатковуваний дохід за день" — теж ОБЧИСЛЮВАНА (аудит PR #579 знайшов цю; графа4 я знайшла
# сама при виправленні — той самий клас, мій попередній коментар про "3=до перерахування"
# був написаний з пам'яті, не звірений живо). Обчислювані графи не входять: сума операції з
# банку могла б випадково збігтися з денним НЕТТО-підсумком і хибно позначитись "уже в книзі".
# Зараз обидві графи 100% порожні в реальній книзі (0/101 заповнених рядків) — практичного
# ризику сьогодні нема, але структурно вони не мали бути в цьому наборі.
_BOOK_MONEY_COLS = (1, 2, 5, 6, 7, 8, 9)
_NO_TELEGRAM = os.environ.get("AUDIT_NO_TELEGRAM") == "1"

_DATE_RE = re.compile(r"(\d{2})\.(\d{2})\.(\d{4})")
_AMOUNT_RE = re.compile(r"^-?[\d\s]+\.\d{2}$")


def _notify(msg: str) -> None:
    if _NO_TELEGRAM:
        return
    try:
        from telegram_notify import send_telegram_message
        send_telegram_message(msg)
    except Exception as e:  # noqa: BLE001
        print(f"[PrivatStmt] Telegram не надіслано (не критично): {e}", file=sys.stderr)


def _all_statements() -> list:
    """Усі PDF-виписки в теці — обробляємо ВСІ щоразу (без курсора, див. докстрінг модуля)."""
    return sorted(set(glob.glob(str(DOCS_DIR / "*" / "ПриватБанк" / "*.pdf"))))


def extract_table(pdf_path: str) -> list:
    """Тонка обгортка: `extract_tables()`, не `extract_text()` — таблична екстракція зберігає
    межі клітинок, текстовий потік для широкої багаторядкової таблиці плутає колонки (звірено
    живо на реальній виписці 2026-09-19).

    Файл — CAdES/PKCS7-підписаний (ASN.1-обгортка ПЕРЕД `%PDF-…`, офсет варіюється — звірено
    живо: 72 байти на реальних виписках 18-19.09, але не гарантовано фіксований), тому шукаємо
    маркер і ріжемо звідти, а не довіряємо файлу з байта 0 — pdfplumber інакше падає
    з "No /Root object"."""
    import pdfplumber
    raw = Path(pdf_path).read_bytes()
    idx = raw.find(b"%PDF-")
    if idx == -1:
        raise ValueError(f"{pdf_path}: немає %PDF- маркера — не PDF (протухле посилання/HTML?)")
    import io
    with pdfplumber.open(io.BytesIO(raw[idx:])) as pdf:
        tables = []
        for page in pdf.pages:
            tables += page.extract_tables()
    rows = []
    for t in tables:
        rows.extend(t)
    return rows


def parse_transactions(table: list) -> list:
    """Чиста функція (тестована без PDF): рядок таблиці — операція, якщо колонка 1 (дата) має
    dd.mm.yyyy, і колонка 3 (сума) парситься як число. Решта (заголовки/підсумки) відсіюється
    природно — вони не мають ОБОХ ознак одночасно."""
    out = []
    for row in table:
        if len(row) < 5:
            continue
        date_cell = (row[1] or "").replace("\n", " ")
        amount_cell = (row[3] or "").replace("\n", " ").strip()
        m = _DATE_RE.search(date_cell)
        if not m or not _AMOUNT_RE.match(amount_cell):
            continue
        try:
            amount = float(amount_cell.replace(" ", ""))
        except ValueError:
            continue
        d, mo, y = m.groups()
        ref = (row[0] or "").replace("\n", " ").strip()
        purpose = (row[4] or "").replace("\n", " ").strip() if len(row) > 4 else ""
        counterparty = (row[6] or "").replace("\n", " ").strip() if len(row) > 6 else ""
        out.append({
            "ref": ref,
            "date": f"{y}-{mo}-{d}",
            "amount": round(amount, 2),
            "purpose": purpose,
            "counterparty": counterparty,
        })
    return out


def _book_money_index() -> dict:
    """READ-ONLY: {сума(грн, 2 знаки, БЕЗ знаку) → [дата рядка, ...]} з БУДЬ-ЯКОЇ грошової графи
    книги (не лише графа2, на відміну від checkbox_registry_sync — виписка покриває всі типи
    рухів). Книгу НЕ пише."""
    index: dict = {}
    if not KODV_XLSX.exists():
        return index
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(KODV_XLSX), data_only=True, read_only=True)
        try:
            ws = wb["КОДВ"]
            for row in ws.iter_rows(min_row=7):
                if len(row) <= max(_BOOK_MONEY_COLS):
                    continue
                d = row[0].value
                for col in _BOOK_MONEY_COLS:
                    v = row[col].value
                    if isinstance(v, (int, float)) and v:
                        key = round(abs(float(v)), 2)
                        index.setdefault(key, []).append(_coerce_date(d))
        finally:
            wb.close()
    except Exception as e:  # noqa: BLE001
        print(f"[PrivatStmt] книжковий індекс не побудовано (не критично): {e}", file=sys.stderr)
    return index


def _coerce_date(v):
    if isinstance(v, datetime):
        return v.date()
    if hasattr(v, "year") and hasattr(v, "month"):
        return v
    return None


def _already_in_book(txn: dict, index: dict) -> bool:
    key = round(abs(txn["amount"]), 2)
    dates = index.get(key, [])
    try:
        txn_date = datetime.strptime(txn["date"], "%Y-%m-%d").date()
    except ValueError:
        return False
    return any(d is not None and abs((d - txn_date).days) <= 1 for d in dates)


def main() -> None:
    # Детектор тиші (аудит Д3, 2026-09-18): ПриватБанк мовчав з 11.08 і про це не було жодного
    # сигналу — той самий клас, що RozetkaPay мав до PR #566. Лист приходить ЩОДНЯ (підтверджено
    # живо 2026-09-19), тому поріг вужчий за RozetkaPay.
    source_freshness.check_and_record(
        "ПриватБанк", str(DOCS_DIR / "*" / "ПриватБанк" / "*.pdf"), max_stale_days=2)
    source_freshness.write_report()

    files = _all_statements()
    if not files:
        print("[PrivatStmt] жодної виписки в документи_КОДВ/*/ПриватБанк/*.pdf ще немає.")
        return

    index = _book_money_index()
    all_txns = []
    any_file_failed = False
    for path in files:
        try:
            table = extract_table(path)
        except Exception as e:  # noqa: BLE001
            print(f"[PrivatStmt] {path}: не вдалось прочитати ({e})", file=sys.stderr)
            _notify(f"🚨 privat_statement_kandydaty: {os.path.basename(path)} не читається: {e}")
            any_file_failed = True
            continue
        for t in parse_transactions(table):
            t["file"] = os.path.basename(path)
            t["in_book"] = _already_in_book(t, index)
            all_txns.append(t)

    # Дедуп intra-run (та сама операція теоретично може повторитись у двох PDF — напр. якщо
    # виписку перезапросили за той самий день): за (ref, дата, сума).
    seen_this_run = set()
    txns = []
    for t in all_txns:
        key = (t["ref"], t["date"], t["amount"])
        if key in seen_this_run:
            continue
        seen_this_run.add(key)
        txns.append(t)

    unresolved = [
        {
            "key": f"{t['ref']}_{t['date']}_{t['amount']}",
            "summary": f"{t['date']} {t['amount']:+.2f}₴ {t['purpose'][:80]}",
            "sum": t["amount"],
            "date": t["date"],
        }
        for t in txns if not t["in_book"]
    ]
    # resolve=not any_file_failed: якщо бодай один PDF не прочитався, txns НЕ гарантовано повний
    # (той самий принцип, що toysi_returns_kandydaty.py) — не закриваємо кандидатів наосліп.
    sync_result = kandydaty_registry.sync_open_candidates(
        "privat_statement", unresolved, resolve=not any_file_failed)
    if sync_result["newly_opened"] or sync_result["resolved"]:
        print(f"[PrivatStmt] Реєстр: +{len(sync_result['newly_opened'])} нових, "
              f"-{len(sync_result['resolved'])} закритих, {len(sync_result['still_open'])} досі відкриті.")
    kandydaty_registry.write_open_report()

    today = datetime.now().strftime("%Y-%m-%d")
    month_dir = DOCS_DIR / datetime.now().strftime("%Y-%m") / "ПриватБанк"
    month_dir.mkdir(parents=True, exist_ok=True)
    (month_dir / f"{today}_privat_kandydaty.json").write_text(
        json.dumps(txns, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    not_in_book = [t for t in txns if not t["in_book"]]
    lines = [
        f"# ПриватБанк — операції з виписки, яких не видно в книзі, {today}",
        "",
        "«У книзі» — сума операції знайдена в БУДЬ-ЯКІЙ графі сирих операцій книги (2,3,6,7,8,9,10;",
        "БЕЗ графи 4/11 — це обчислювані підсумки, не сирі операції) на ту саму дату ±1 день.",
        "Кирилична деталізація (призначення/контрагент) у файлі кандидатів (json поруч), сюди не",
        "виносимо — короткий огляд.",
        "",
        "| Дата | Сума | Документ | У книзі? |",
        "|---|---|---|---|",
    ]
    for t in txns:
        mark = "✅" if t["in_book"] else "❗ НЕ в книзі"
        lines.append(f"| {t['date']} | {t['amount']:+.2f} | {t['ref'] or '—'} | {mark} |")
    (month_dir / f"{today}_privat_kandydaty.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    print(f"[PrivatStmt] {len(txns)} операцій усього, {len(not_in_book)} не в книзі.")
    if any_file_failed:
        print("[PrivatStmt] ⚠️ бодай один PDF не прочитався — закриття кандидатів пропущено цим прогоном.")


if __name__ == "__main__":
    main()
