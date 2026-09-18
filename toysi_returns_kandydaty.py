"""
toysi_returns_kandydaty.py — кандидати повернень Toysi з таблиці «Взаєморозрахунки»
(https://toysi.ua/contact_info/?deposit=).

НАВІЩО (незалежний аудитор, КОДВ_журнал.md «ДОПОВНЕННЯ 5», Priority #3, 2026-09-18):
повернення — ЄДИНИЙ клас кандидатів автоматики КОДВ із влучністю 0/7. Ніхто не читає
«Взаєморозрахунки» автоматично — 7 реальних повернень (1 709,96₴ обороту) знайдено лише
розовим ручним аудитом. Одне з них (Куровська, toysi-100449926, ТС000001089) висіло
НЕВІДОМИМ книзі 20 днів, поки вручну не спіймали.

ЩО РОБИТЬ: заходить у кабінет Toysi (та сама storageState-сесія, що toysi_cabinet_scraper.py),
читає «Взаєморозрахунки» — сторінка сама показує посилання на завантаження звіту за кожен
напівмісяць (динамічно, контрагент-ID НЕ хардкодиться — береться з посилань на сторінці).
Завантажує ВСІ доступні періоди (Playwright APIRequestContext, авторизований GET — НЕ
`page.goto()`, той кидає виняток на прямому завантаженні файлу), знаходить рядки з
«Повернення товарів» у графі «Документ».

ЧОМУ НЕ РАХУЄ СУМУ АВТОМАТИЧНО (звірено живо структуру книги, рядок 106): бухгалтер
записує повернення НЕ як пряме дзеркало суми з Toysi — графа 9 книги містить ЧИСТУ ВТРАТУ
(собівартість − реально повернене Toysi), а повний наратив (toysi-id, ТС-номер, дати,
розрахунок) — у графі 5 текстом. Це судження бухгалтера (яка частка повернення — наша
втрата, яка — Toysi), не механічний перенос. Тому скрипт лише ЗНАХОДИТЬ факт повернення
в Toysi і звіряє ТЕКСТОВИМ пошуком (toysi-id АБО ТС-номер у графі 5 будь-якого рядка
книги), чи бухгалтер про нього вже знає — суму й трактування рахує бухгалтер сама.

Крос-звірка з книгою — READ-ONLY. Пише кандидатів у документи_КОДВ/YYYY-MM/Toysi/, книгу
НЕ чіпає. Реєстр відкритих кандидатів — той самий kandydaty_registry.py, що вже закрив
Д1/Д4 для Checkbox (PR #564): курсор джерела ≠ курсор книги, кандидат не губиться, поки
бухгалтер його дійсно не внесе.

ЗАПУСК: python toysi_returns_kandydaty.py
"""
import json
import os
import re
import sys
from datetime import date, datetime
from pathlib import Path

from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import kandydaty_registry

BASE_DIR = Path(__file__).parent
COWORK_DIR = Path(os.environ.get(
    "PLUTUS_COWORK_DIR", r"C:\Users\smach\Claude\Projects\PlutusToys_avtonomiya"))
KODV_XLSX = COWORK_DIR / "KODV_PlutusToys_2026.xlsx"
DOCS_DIR = COWORK_DIR / "документи_КОДВ"

STATE_FILE = Path(
    os.environ.get("TOYSI_CABINET_STATE_FILE", str(BASE_DIR / ".local_secrets" / "toysi_cabinet_state.json")))

SETTLEMENTS_URL = "https://toysi.ua/contact_info/?deposit="
NAV_TIMEOUT_MS = 30000
_NO_TELEGRAM = os.environ.get("AUDIT_NO_TELEGRAM") == "1"

_RETURN_MARKER = "повернення товарів"
_TS_RE = re.compile(r"ТС\d+", re.IGNORECASE)
_TOYSI_ORDER_RE = re.compile(r"toysi-(\d+)", re.IGNORECASE)


class ToysiReturnsError(Exception):
    pass


def _log(msg: str) -> None:
    print(f"[ToysiReturns] {msg}")


def _notify(msg: str) -> None:
    if _NO_TELEGRAM:
        return
    try:
        from telegram_notify import send_telegram_message
        send_telegram_message(msg)
    except Exception as e:  # noqa: BLE001
        print(f"[ToysiReturns] Telegram не надіслано (не критично): {e}", file=sys.stderr)


def _discover_period_links(page) -> list:
    """Сторінка «Взаєморозрахунки» САМА показує посилання на завантаження за доступні
    напівмісяці — контрагент-ID у ?part=<ID>_<MM>_<YYYY>_partN беремо звідси, не
    хардкодимо (ID міг би змінитись, і хардкод про це б мовчав)."""
    hrefs = page.eval_on_selector_all(
        "a[href*='action=get_history']", "els => els.map(e => e.href)")
    return sorted(set(hrefs))


def fetch_return_rows() -> list:
    """Заходить у кабінет, читає ВСІ доступні періоди «Взаєморозрахунки», повертає рядки
    з «Повернення товарів» у графі «Документ»: [{"doc": str, "toysi_order_id": int|None,
    "tc_number": str|None, "date": "YYYY-MM-DD"|None, "sum_debet": float|None, "period": str}]."""
    if not STATE_FILE.exists():
        raise ToysiReturnsError(f"нема сесії ({STATE_FILE.name}) — `python toysi_cabinet_scraper.py --login`")

    import openpyxl
    from io import BytesIO

    rows_out = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        ctx = browser.new_context(storage_state=str(STATE_FILE))
        page = ctx.new_page()
        try:
            page.goto(SETTLEMENTS_URL, timeout=NAV_TIMEOUT_MS, wait_until="domcontentloaded")
            page.wait_for_timeout(2000)
            if "auth" in page.url.lower() or "login" in page.url.lower():
                raise ToysiReturnsError(
                    f"сесію НЕ прийнято: опинились на {page.url} (треба --login)")
            links = _discover_period_links(page)
            if not links:
                raise ToysiReturnsError(
                    "жодного посилання на звіт не знайдено на сторінці «Взаєморозрахунки» — "
                    "могла змінитись верстка, перевір вручну")
            _log(f"знайдено {len(links)} періодів на сторінці.")

            for url in links:
                m = re.search(r"part=(\d+_\d{2}_\d{4}_part\d)", url)
                period = m.group(1) if m else url
                try:
                    resp = page.request.get(url, timeout=NAV_TIMEOUT_MS)
                    if resp.status != 200:
                        _log(f"⚠️ {period}: HTTP {resp.status} — пропускаю.")
                        continue
                    wb = openpyxl.load_workbook(BytesIO(resp.body()), data_only=True)
                except Exception as e:  # noqa: BLE001 — один період не має валити решту
                    _log(f"⚠️ {period}: не вдалось прочитати ({e}) — пропускаю.")
                    continue
                ws = wb.active
                for row in ws.iter_rows(min_row=6, values_only=True):
                    if not row or not row[0]:
                        continue
                    doc = str(row[0])
                    if _RETURN_MARKER not in doc.lower():
                        continue
                    tc_m = _TS_RE.search(doc)
                    rows_out.append({
                        "doc": doc,
                        "toysi_order_id": row[1] if len(row) > 1 else None,
                        "tc_number": tc_m.group(0) if tc_m else None,
                        "date": str(row[2]) if len(row) > 2 and row[2] else None,
                        "sum_debet": row[6] if len(row) > 6 else None,
                        "period": period,
                    })
        finally:
            browser.close()
    return rows_out


def _book_narrative_text() -> str:
    """READ-ONLY: увесь текст графи 5 (наратив/опис) книги, злитий в один рядок для
    підрядкового пошуку toysi-id/ТС-номера. Книгу не пише, не парсить структуровано —
    повернення записуються текстом, не окремою колонкою (звірено живо, рядок 106)."""
    if not KODV_XLSX.exists():
        return ""
    import openpyxl
    wb = openpyxl.load_workbook(str(KODV_XLSX), data_only=True, read_only=True)
    try:
        ws = wb["КОДВ"]
        parts = []
        for row in ws.iter_rows(min_row=7):
            if len(row) < 5:
                continue
            v = row[4].value  # графа 5 — опис/наратив
            if v:
                parts.append(str(v))
        return " \n".join(parts)
    finally:
        wb.close()


def _already_in_book(row: dict, book_text: str) -> bool:
    """Чи згадано toysi-id АБО ТС-номер цього повернення хоч у ОДНОМУ рядку книги.
    Текстовий підрядковий пошук, не структурований збіг — так само, як бухгалтер сама
    записує ці рядки (наратив, не колонка)."""
    if row.get("toysi_order_id") and f"toysi-{row['toysi_order_id']}" in book_text:
        return True
    if row.get("tc_number") and row["tc_number"] in book_text:
        return True
    return False


def main() -> None:
    try:
        rows = fetch_return_rows()
    except ToysiReturnsError as e:
        _log(f"помилка: {e}")
        _notify(f"🚨 toysi_returns_kandydaty: {e}")
        sys.exit(1)
    except Exception as e:  # noqa: BLE001
        _log(f"несподівана помилка: {e}")
        _notify(f"🚨 toysi_returns_kandydaty: несподівана помилка: {e}")
        sys.exit(1)

    if not rows:
        _log("Повернень у доступних періодах не знайдено.")
        return

    book_text = _book_narrative_text()
    for r in rows:
        r["in_book"] = _already_in_book(r, book_text)

    unresolved = [
        {
            "key": r["tc_number"] or f"toysi-{r['toysi_order_id']}" or r["doc"][:40],
            "summary": f"{r['doc'][:100]} (сума дебет {r.get('sum_debet')})",
            "sum": r.get("sum_debet"),
            "date": r.get("date"),
        }
        for r in rows if not r["in_book"]
    ]
    sync_result = kandydaty_registry.sync_open_candidates("toysi_returns", unresolved)
    if sync_result["newly_opened"] or sync_result["resolved"]:
        _log(f"Реєстр відкритих кандидатів: +{len(sync_result['newly_opened'])} нових, "
             f"-{len(sync_result['resolved'])} закритих, {len(sync_result['still_open'])} досі відкриті.")
    kandydaty_registry.write_open_report()

    today = datetime.now().strftime("%Y-%m-%d")
    month_dir = DOCS_DIR / datetime.now().strftime("%Y-%m") / "Toysi"
    month_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{today}_toysi_povernennya_kandydaty"
    (month_dir / f"{stem}.json").write_text(
        json.dumps(rows, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    lines = [
        f"# Toysi — повернення товарів, знайдені у «Взаєморозрахунках», {today}",
        "",
        "Джерело: https://toysi.ua/contact_info/?deposit= (усі доступні періоди). «У книзі»",
        "— toysi-id або ТС-номер згадано хоч в одному рядку графи 5 книги (текстовий пошук,",
        "не структурований збіг — повернення записуються наративом, не окремою колонкою).",
        "",
        "| Документ | Toysi ID | ТС-номер | Дата | Сума дебет | У книзі? |",
        "|---|---|---|---|---|---|",
    ]
    for r in rows:
        mark = "✅" if r["in_book"] else "❗ НЕ в книзі"
        lines.append(f"| {r['doc'][:80]} | {r.get('toysi_order_id', '?')} | "
                     f"{r.get('tc_number', '?')} | {r.get('date', '?')} | "
                     f"{r.get('sum_debet', '?')} | {mark} |")
    (month_dir / f"{stem}.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    n_open = len(unresolved)
    _log(f"ГОТОВО: {len(rows)} повернень знайдено ({n_open} ще не в книзі) → {month_dir / f'{stem}.md'}")
    if n_open:
        _notify(f"↩️ Toysi: {n_open} повернень ще НЕ в книзі (з {len(rows)} знайдених у "
                f"«Взаєморозрахунках»). Див. {stem}.md")


if __name__ == "__main__":
    main()
