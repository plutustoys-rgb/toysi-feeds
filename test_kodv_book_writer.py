#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
test_kodv_book_writer.py — регрес-тест динамічної межі шаблону книги КОДВ.

Живий інцидент (Аудитор, КОДВ_CHANNEL.md, 2026-09-28): форму книги розширили
25.09 з рядка 126 до 300 (РАЗОМ переїхав у рядок 301, SUM(...7:300)), але
хардкод-константу `LAST_TEMPLATE_ROW = 126` не оновили. `append_row` відмовлявся
писати рядки 127+ ("за межами підготовленого шаблону, до 126") — рядки 127-134
довелось внести ВРУЧНУ, в обхід механічної перевірки дублів цього ж модуля
(саме той захист, заради якого модуль існує).

ФІКС: `_find_template_last_row(ws)` шукає рядок «РАЗОМ» у стовпці A ЖИВО при
кожному відкритті книги, замість хардкод-константи — форму можна розширювати
без правки коду.

ІНВАРІАНТИ:
  1. Межа шаблону визначається за позицією «РАЗОМ», не константою.
  2. Розширення форми (РАЗОМ переїхав далі) підхоплюється АВТОМАТИЧНО, без
     редеплою коду — головне, що це фіксить.
  3. Рядок за межею (target_row > межа) — RuntimeError, як і раніше (безпека
     не втрачена, лише межа тепер жива).
  4. Відсутність рядка «РАЗОМ» у книзі — явний RuntimeError, не мовчазний
     здогад/крах на IndexError.

Синтетична тимчасова книга (той самий мінімальний набір колонок, що й
`_DATA_COLS`), справжню книгу НЕ чіпає. `python test_kodv_book_writer.py` → exit 0/1.
"""
import sys
import tempfile
from pathlib import Path

import openpyxl

import kodv_book_writer as kbw

_FAILS = []


def _chk(name, cond):
    print(f"[{'OK ' if cond else 'FAIL'}] {name}")
    if not cond:
        _FAILS.append(name)


def _make_book(tmpdir, last_data_row: int, razom_row: int) -> Path:
    """Мінімальна синтетична книга: рядки 7..last_data_row заповнені (Графа 5 = унікальний
    текст без номера документа, щоб не зачепити перевірку дублів), 「РАЗОМ」 у razom_row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = kbw.SHEET_NAME
    for r in range(kbw.FIRST_DATA_ROW, last_data_row + 1):
        ws.cell(r, kbw.COL_DATE, "2026-09-01")
        ws.cell(r, kbw.COL_INCOME, 100)
        ws.cell(r, kbw.COL_DOC, f"тестовий запис рядка {r}")
    ws.cell(razom_row, kbw.COL_DATE, "РАЗОМ")
    path = Path(tmpdir) / "test_book.xlsx"
    wb.save(path)
    return path


# ── 1+2: межа шаблону — за позицією РАЗОМ, не константою ──
with tempfile.TemporaryDirectory() as tmpdir:
    book = _make_book(tmpdir, last_data_row=10, razom_row=16)
    wb = openpyxl.load_workbook(book, data_only=False)
    ws = wb[kbw.SHEET_NAME]
    _chk("межа шаблону = razom_row - 1 (16-1=15)", kbw._find_template_last_row(ws) == 15)
    _chk("остання заповнена = 10", kbw._find_last_data_row(ws) == 10)

# 2б: та сама логіка, РАЗОМ переїхав ДАЛІ (розширення форми) — підхоплюється БЕЗ правки коду
with tempfile.TemporaryDirectory() as tmpdir:
    book = _make_book(tmpdir, last_data_row=10, razom_row=301)
    wb = openpyxl.load_workbook(book, data_only=False)
    ws = wb[kbw.SHEET_NAME]
    _chk("розширена форма: межа = 300 (той самий живий випадок, що в реальній книзі)",
         kbw._find_template_last_row(ws) == 300)

# ── ПРЕ-ФІКС інваріант: живий інцидент дослівно (126→300, дані до 126) ──
# Стара LAST_TEMPLATE_ROW=126 заблокувала б append на рядку 127 ("за межами шаблону, до 126"),
# хоча РАЗОМ насправді на 301 (межа 300) — саме цей інцидент фікс закриває.
with tempfile.TemporaryDirectory() as tmpdir:
    book = _make_book(tmpdir, last_data_row=126, razom_row=301)
    report = kbw.append_row(
        date="2026-09-28", graph5="тестовий рахунок №99999", graph9=10.0,
        book_path=book, dry_run=True,
    )
    _chk("живий інцидент: рядок 127 (за старою межею 126, у межах живої 300) ЗАПИСУЄТЬСЯ",
         report["row"] == 127)

# ── 3: рядок ЗА живою межею — RuntimeError, безпека не втрачена ──
with tempfile.TemporaryDirectory() as tmpdir:
    book = _make_book(tmpdir, last_data_row=15, razom_row=16)  # межа=15, дані вже до 15 — рядка 16 нема
    try:
        kbw.append_row(date="2026-09-28", graph5="тест", graph9=10.0, book_path=book, dry_run=True)
        _chk("за межею шаблону: RuntimeError", False)
    except RuntimeError as e:
        _chk("за межею шаблону: RuntimeError з поясненням", "за межами підготовленого шаблону" in str(e))

# ── 4: РАЗОМ відсутній узагалі — явний RuntimeError, не мовчазний здогад/крах ──
with tempfile.TemporaryDirectory() as tmpdir:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = kbw.SHEET_NAME
    ws.cell(kbw.FIRST_DATA_ROW, kbw.COL_DATE, "2026-09-01")  # дані є, РАЗОМ нема
    path = Path(tmpdir) / "no_razom.xlsx"
    wb.save(path)
    wb2 = openpyxl.load_workbook(path, data_only=False)
    ws2 = wb2[kbw.SHEET_NAME]
    try:
        kbw._find_template_last_row(ws2)
        _chk("без РАЗОМ: RuntimeError", False)
    except RuntimeError as e:
        _chk("без РАЗОМ: явний RuntimeError, не крах", "РАЗОМ" in str(e))


if _FAILS:
    print(f"\n❌ ПРОВАЛЕНО: {len(_FAILS)} — {_FAILS}")
    sys.exit(1)
print("\n✅ Межа шаблону книги КОДВ визначається живо за рядком «РАЗОМ», "
      "не застарілою константою (аудит 2026-09-28)")
sys.exit(0)
