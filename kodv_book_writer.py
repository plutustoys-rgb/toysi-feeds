"""kodv_book_writer.py — єдина точка запису НОВОГО рядка в книгу КОДВ
(`KODV_PlutusToys_2026.xlsx`), з механічною перевіркою дубля номера документа
в Графі 5 ПЕРЕД записом — у коді, не в пам'яті агента.

ЧОМУ (архітектурний аудит 2026-09-18): рядок 80 книги (акт Prom №11865348)
послався на "Рахунок на оплату покупцю №12128253-66", який уже повністю
визнаний витратою в рядку 7 ("рахунок №UA-12128253-66" — той самий номер,
інший префікс). Подвійний облік 734.36 грн. КОДВ перевірила арифметику
ВСЕРЕДИНІ акту (порахувала правильно) і на цьому зупинилась — не
перевірила, чи номер документа вже є деінде в книзі. Знайшов це лише
незалежний Agent-аудит, звіряючи номер документа, не суму. Корінь
(див. пам'ять kodv-confirmatory-vs-adversarial-checking): самоперевірка
ПІДТВЕРДЖУВАЛЬНА (перевіряє те, у чому засумнівались), аудит ЗМАГАЛЬНИЙ
(перевіряє все). Жодна кількість написаних правил це не лікує, бо агент
не бачить того, у чому не сумнівається — тому перевірку винесено в код,
що виконується щоразу, незалежно від того, чи агент здогадався її
запустити.

НЕ ЗАМІНЮЄ незалежний Agent-аудит (§6.3 брифу) чи тижневий scheduled
аудит (`kodv-independent-book-audit`) — це дешевий детермінований
ПЕРШИЙ фільтр САМЕ для дубля номера документа. Логічні помилки (чи це
справді нова витрата, чи вже врахована іншим способом; чи розподіл
пропорційний коректний тощо) лишаються за живим аудитом.

СТРУКТУРА КНИГИ (типова форма, Наказ Мінфіну 13.05.2021 №261), звірено
живо (не вигадано): аркуш "КОДВ", рядок 5 = підписи "Графа N", рядок 6 =
описові заголовки, дані з рядка 7. Стовпці A-K = Графи 1-11 офіційної
форми (Графа 5 = "Реквізити підтвердного документа", Графа 9 = "Інші
витрати"); стовпець L = робочі примітки/історія аудиту, НЕ частина
типової форми. Графи 4 (D) і 11 (K) — ФОРМУЛИ (`=IF(B{r}="","",...)`),
ПРЕ-ЗАПОВНЕНІ на рядки наперед (перевірено: рядок 106 — останній із
датою, рядки 107-126 уже мають формули, порожні лише A/B/C/E/F/G/H/I/J).
Тому НОВИЙ рядок — це ЗАПОВНЕННЯ наступного порожнього рядка з формулами,
НЕ вставка/зсув рядків (вставка зламала б `РАЗОМ` у рядку 128:
`=SUM(B7:B126)` не адаптується автоматично при `ws.insert_rows`).

ЗАПУСК:
  python kodv_book_writer.py append --date 2026-09-18 --graph9 123.45 \
      --graph5 "Опис + номер документа" [--graph2 N] [--note "..."] \
      [--confirm-duplicate "чому це НЕ дублікат"] [--dry-run]
  python kodv_book_writer.py check "текст з номером документа, який плануєш внести"
  python kodv_book_writer.py scan   # усі номери документів книги, що зустрічаються >1 разу
"""
import argparse
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import openpyxl  # noqa: E402

BOOK_PATH = Path(r"C:\Users\smach\Claude\Projects\PlutusToys_avtonomiya\KODV_PlutusToys_2026.xlsx")
SHEET_NAME = "КОДВ"
FIRST_DATA_ROW = 7
LAST_TEMPLATE_ROW = 126   # РАЗОМ (128) рахує SUM(...7:126) — за межі не виходимо без ручного розширення форми
COL_DATE, COL_INCOME, COL_RETURNS, COL_DOC, COL_COGS = 1, 2, 3, 5, 6
COL_LABOR, COL_TAXES, COL_OTHER_EXP, COL_AMORT, COL_NOTE = 7, 8, 9, 10, 12

# Номер документа завжди йде після "№" (українська типографська норма рахунків/актів/ТТН);
# мінімум 3 символи після символу, щоб не ловити шум. Порівнюємо НЕ сирий рядок, а нормалізовану
# цифрову форму — інакше "№UA-12128253-66" і "№12128253-66" (той самий рахунок, різний префікс,
# РЕАЛЬНИЙ кейс інциденту) не збіглися б при точному порівнянні рядків.
_DOC_NO_RE = re.compile(r"№\s*([A-Za-zА-ЯҐЄІЇа-яґєії0-9][\w\-/]{2,})", re.UNICODE)
_MIN_DIGITS_FOR_NUMERIC_KEY = 5
# "реєстр №..." — номер ПЛАТІЖНОГО БАТЧА (NovaPay/банк), не унікального документа-витрати:
# один реєстр законно покриває кілька РІЗНИХ замовлень різних площадок в один день (перевірено
# живо на реальних рядках 63/68 книги — той самий реєстр №15991878, різні EVA/Rozetka замовлення,
# не подвійний облік). Виключаємо з перевірки дублів, інакше щоденний false positive на кожен
# спільний реєстр приховає РЕАЛЬНІ збіги серед шуму.
_REGISTER_PREFIX_RE = re.compile(r"реєстр(?:у|ом)?\s*$", re.IGNORECASE | re.UNICODE)


class DuplicateDocumentReferenceError(Exception):
    def __init__(self, matches):
        self.matches = matches
        lines = [
            f"  рядок {m['row']}: номер «{m['existing_raw']}» (новий: «{m['new_raw']}»)"
            for m in matches
        ]
        super().__init__(
            "МОЖЛИВИЙ ПОДВІЙНИЙ ОБЛІК — цей номер документа вже фігурує в книзі:\n"
            + "\n".join(lines)
            + "\n\nЯкщо це справді НЕ дублікат (інша послуга того самого контрагента, "
              "часткове списання тощо) — повтори виклик з --confirm-duplicate \"причина\"."
        )


def _normalize(token: str) -> str:
    digits = re.sub(r"\D", "", token)
    if len(digits) >= _MIN_DIGITS_FOR_NUMERIC_KEY:
        return digits
    return token.strip().casefold()


def extract_doc_numbers(text: str) -> dict:
    """raw-токен (як у тексті) -> нормалізований ключ для порівняння.
    Пропускає "реєстр №..." (платіжний батч, легітимно спільний для кількох рядків)."""
    if not text:
        return {}
    out = {}
    for m in _DOC_NO_RE.finditer(text):
        prefix = text[max(0, m.start() - 12): m.start()]
        if _REGISTER_PREFIX_RE.search(prefix):
            continue
        out[m.group(1)] = _normalize(m.group(1))
    return out


def _find_last_data_row(ws) -> int:
    last = FIRST_DATA_ROW - 1
    for r in range(FIRST_DATA_ROW, LAST_TEMPLATE_ROW + 1):
        if ws.cell(r, COL_DATE).value is not None:
            last = r
    return last


def build_doc_index(ws, up_to_row: int) -> dict:
    """нормалізований_ключ -> [(рядок, сирий_токен), ...] для всіх рядків 7..up_to_row."""
    index: dict = {}
    for r in range(FIRST_DATA_ROW, up_to_row + 1):
        text = ws.cell(r, COL_DOC).value
        for raw, key in extract_doc_numbers(text or "").items():
            index.setdefault(key, []).append((r, raw))
    return index


def find_duplicates(new_text: str, index: dict) -> list:
    matches = []
    for new_raw, key in extract_doc_numbers(new_text).items():
        for row, existing_raw in index.get(key, []):
            matches.append({"row": row, "existing_raw": existing_raw, "new_raw": new_raw})
    return matches


def check_duplicates(text: str, book_path: Path = BOOK_PATH) -> list:
    wb = openpyxl.load_workbook(book_path, data_only=False)
    ws = wb[SHEET_NAME]
    last_row = _find_last_data_row(ws)
    index = build_doc_index(ws, last_row)
    return find_duplicates(text, index)


def scan_book_for_internal_duplicates(book_path: Path = BOOK_PATH) -> dict:
    """Номери документів, що зустрічаються В КНИЗІ вже зараз більш ніж в одному рядку."""
    wb = openpyxl.load_workbook(book_path, data_only=False)
    ws = wb[SHEET_NAME]
    last_row = _find_last_data_row(ws)
    index = build_doc_index(ws, last_row)
    return {key: rows for key, rows in index.items() if len(rows) > 1}


def append_row(
    *,
    date,
    graph5: str,
    graph2=None, graph3=None, graph6=None, graph7=None, graph8=None,
    graph9=None, graph10=None,
    note: str | None = None,
    confirm_duplicate_reason: str | None = None,
    dry_run: bool = False,
    book_path: Path = BOOK_PATH,
) -> dict:
    if not graph5 or not graph5.strip():
        raise ValueError("graph5 (Реквізити підтвердного документа) обов'язковий — без опису немає за чим звіряти дублі.")
    wb = openpyxl.load_workbook(book_path, data_only=False)
    ws = wb[SHEET_NAME]
    last_row = _find_last_data_row(ws)
    target_row = last_row + 1
    if target_row > LAST_TEMPLATE_ROW:
        raise RuntimeError(
            f"Рядок {target_row} за межами підготовленого шаблону (до {LAST_TEMPLATE_ROW}). "
            "Форму треба розширити вручну (формули Графи 4/11 + діапазон РАЗОМ) — не пишу наосліп."
        )
    if ws.cell(target_row, COL_DATE).value is not None:
        raise RuntimeError(
            f"Рядок {target_row} вже має дані в Графі 1 — розбіжність із очікуваним шаблоном, "
            "не пишу без ручної перевірки книги."
        )

    index = build_doc_index(ws, last_row)
    dups = find_duplicates(graph5, index)
    if dups and not confirm_duplicate_reason:
        raise DuplicateDocumentReferenceError(dups)

    report = {
        "row": target_row,
        "duplicates_found": dups,
        "confirmed_reason": confirm_duplicate_reason,
        "dry_run": dry_run,
    }
    if dry_run:
        return report

    backup_path = book_path.with_name(
        f"{book_path.stem}_backup_{datetime.now():%Y-%m-%d_%H%M%S}_pre-row{target_row}{book_path.suffix}"
    )
    shutil.copy2(book_path, backup_path)
    report["backup"] = str(backup_path)

    values = {
        COL_DATE: date, COL_INCOME: graph2, COL_RETURNS: graph3, COL_DOC: graph5,
        COL_COGS: graph6, COL_LABOR: graph7, COL_TAXES: graph8, COL_OTHER_EXP: graph9,
        COL_AMORT: graph10,
    }
    for col, val in values.items():
        if val is not None:
            ws.cell(target_row, col, val)
    if note:
        ws.cell(target_row, COL_NOTE, note)
    if confirm_duplicate_reason:
        existing_note = ws.cell(target_row, COL_NOTE).value or ""
        ws.cell(
            target_row, COL_NOTE,
            (existing_note + "\n" if existing_note else "")
            + f"⚠️ Підтверджений НЕ-дублікат номера документа: {confirm_duplicate_reason}",
        )
    wb.save(book_path)
    return report


def _cli() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    p_check = sub.add_parser("check", help="Перевірити текст на дублі номера документа, нічого не пишучи.")
    p_check.add_argument("text", help="Текст графи 5, який плануєш внести.")

    p_scan = sub.add_parser("scan", help="Знайти номери документів, що вже зустрічаються в книзі >1 разу.")

    p_append = sub.add_parser("append", help="Дописати новий рядок у перший вільний підготовлений рядок.")
    p_append.add_argument("--date", required=True, help="YYYY-MM-DD")
    p_append.add_argument("--graph5", required=True, help="Реквізити підтвердного документа (обов'язково)")
    p_append.add_argument("--graph2", type=float, default=None, help="Сума доходу за день")
    p_append.add_argument("--graph3", type=float, default=None, help="Повернення коштів")
    p_append.add_argument("--graph6", type=float, default=None, help="Витрати на придбання товарів")
    p_append.add_argument("--graph7", type=float, default=None, help="Оплата праці/ЦПХ")
    p_append.add_argument("--graph8", type=float, default=None, help="ЄСВ, податки, збори")
    p_append.add_argument("--graph9", type=float, default=None, help="Інші витрати")
    p_append.add_argument("--graph10", type=float, default=None, help="Амортизація")
    p_append.add_argument("--note", default=None, help="Робоча примітка (стовпець L)")
    p_append.add_argument("--confirm-duplicate", dest="confirm_duplicate", default=None,
                           help="Явне пояснення, чому знайдений збіг номера документа — НЕ подвійний облік")
    p_append.add_argument("--dry-run", action="store_true", help="Лише показати, що сталося б, нічого не писати")

    args = ap.parse_args()

    if args.cmd == "check":
        dups = check_duplicates(args.text)
        if not dups:
            print("[kodv_book_writer] Дублів номера документа не знайдено.")
            return 0
        print("[kodv_book_writer] МОЖЛИВІ ДУБЛІ:")
        for m in dups:
            print(f"  рядок {m['row']}: «{m['existing_raw']}» (у тексті: «{m['new_raw']}»)")
        return 3

    if args.cmd == "scan":
        dupes = scan_book_for_internal_duplicates()
        if not dupes:
            print("[kodv_book_writer] У книзі немає номерів документів, що повторюються в кількох рядках.")
            return 0
        print(f"[kodv_book_writer] Знайдено {len(dupes)} номер(ів) документа у кількох рядках:")
        for key, rows in dupes.items():
            where = ", ".join(f"рядок {r} («{raw}»)" for r, raw in rows)
            print(f"  {key}: {where}")
        return 3

    if args.cmd == "append":
        try:
            date_val = datetime.strptime(args.date, "%Y-%m-%d")
        except ValueError:
            print(f"[kodv_book_writer] Невалідна дата: {args.date} (очікую YYYY-MM-DD)", file=sys.stderr)
            return 2
        try:
            report = append_row(
                date=date_val, graph5=args.graph5, graph2=args.graph2, graph3=args.graph3,
                graph6=args.graph6, graph7=args.graph7, graph8=args.graph8,
                graph9=args.graph9, graph10=args.graph10, note=args.note,
                confirm_duplicate_reason=args.confirm_duplicate, dry_run=args.dry_run,
            )
        except DuplicateDocumentReferenceError as e:
            print(f"[kodv_book_writer] ВІДМОВЛЕНО, НІЧОГО НЕ ЗАПИСАНО:\n{e}", file=sys.stderr)
            return 3
        if report["dry_run"]:
            print(f"[kodv_book_writer] DRY-RUN: записав би в рядок {report['row']}. "
                  f"Дублі: {report['duplicates_found'] or 'немає'}.")
        else:
            print(f"[kodv_book_writer] Записано в рядок {report['row']}. Бекап: {report['backup']}.")
            if report["duplicates_found"]:
                print(f"[kodv_book_writer] Підтверджений НЕ-дублікат: {report['confirmed_reason']}")
        return 0
    return 1


if __name__ == "__main__":
    sys.exit(_cli())
